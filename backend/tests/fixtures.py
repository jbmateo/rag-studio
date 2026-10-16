"""Builders for test documents and fake DNS resolvers (no binary fixtures committed)."""
import ipaddress
import socket
from pathlib import Path

import docx
from fpdf import FPDF


def make_pdf(tmp_path: Path, pages: list[str]) -> Path:
    pdf = FPDF()
    pdf.set_font("Helvetica", size=12)
    for text in pages:
        pdf.add_page()
        pdf.multi_cell(0, 10, text)
    path = tmp_path / "doc.pdf"
    pdf.output(str(path))
    return path


def make_docx(tmp_path: Path, sections: list[tuple[str, str]]) -> Path:
    document = docx.Document()
    for heading, body in sections:
        document.add_heading(heading, level=1)
        document.add_paragraph(body)
    path = tmp_path / "doc.docx"
    document.save(str(path))
    return path


def _entries(ip: str, port):
    family = socket.AF_INET6 if ":" in ip else socket.AF_INET
    sockaddr = (ip, port or 0, 0, 0) if family == socket.AF_INET6 else (ip, port or 0)
    return [(family, socket.SOCK_STREAM, 6, "", sockaddr)]


def fake_resolver(mapping: dict[str, str]):
    """A getaddrinfo-compatible resolver: mapped hosts resolve; IP literals return themselves."""

    def resolve(host, port=None, *args, **kwargs):
        try:
            ipaddress.ip_address(host)
            return _entries(host, port)
        except ValueError:
            pass
        if host not in mapping:
            raise socket.gaierror(f"unmapped host {host}")
        return _entries(mapping[host], port)

    return resolve


real_ip_resolver = fake_resolver({})
